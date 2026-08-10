from __future__ import annotations

import json
import os
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace

from dramatiq import Worker
from dramatiq.brokers.stub import StubBroker
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image

from physics_vault_api.config import LessonExportSettings, TaskQueueSettings
from physics_vault_api.repositories.import_tasks import SQLiteImportTaskRepository
from physics_vault_api.routers.lesson_exports import build_lesson_exports_router
from physics_vault_api.routers.tasks import build_tasks_router
from physics_vault_api.schemas.lesson_exports import LessonExportRequest
from physics_vault_api.services.document_pipeline import (
    DocumentCleaningService,
    ImportPipelineService,
    PandocAdapter,
    StructuredQuestionParsingService,
)
from physics_vault_api.services.lesson_exports import LessonExportService, _office_text_parts
from physics_vault_api.services.task_center import TaskCenterService
from physics_vault_api.services.task_queue import LessonExportDispatcher
from physics_vault_api.tasks.lesson_exports import build_lesson_export_actors


def test_docx_text_parts_remove_accidental_blank_lines() -> None:
    parts = _office_text_parts("第一行\n\n\n第二行 $v_1$\r\n\r\n第三行")

    assert parts == [
        ("text", "第一行\n第二行 "),
        ("math", "v_1"),
        ("text", "\n第三行"),
    ]


def _lesson_package(image_path: str) -> dict:
    return {
        "id": "lesson-export-test",
        "title": "机械能守恒练习",
        "subtitle": "学生版与教师版共用快照",
        "source": "compose",
        "questions": [
            {
                "question_id": "q-1",
                "question_type": "single_choice",
                "title": "质量为 $m$ 的小球速度满足 $v^2=2gh$，拉力记为 $F_1$，下列说法正确的是？ ![fig:fig-1]",
                "options": [
                    {"opt": "A", "content": "动能为 $\\frac{1}{2}mv^2$"},
                    {"opt": "B", "content": "重力势能为 $mgh$"},
                    {"opt": "C", "content": "机械能不守恒"},
                    {"opt": "D", "content": "无法判断"},
                ],
                "answer": "A、B",
                "analysis": "由 $\\frac{1}{2}mv^2=mgh$ 可知机械能守恒。",
                "figures": [
                    {
                        "fig_uuid": "fig-1",
                        "local_path": image_path,
                        "display_scale": 55,
                        "display_align": "center",
                        "caption": "小球运动示意图",
                    }
                ],
                "knowledge_point": "机械能守恒",
                "knowledge_points": [],
            }
        ],
        "knowledgeCards": [
            {
                "id": "k-1",
                "title": "机械能守恒",
                "summary": "只有重力做功时，机械能保持不变。",
                "points": ["明确研究对象", "选择零势能面", "列出守恒方程"],
                "relatedQuestionIds": ["q-1"],
            }
        ],
        "textBlocks": [
            {
                "id": "text-title",
                "title": "机械能专题训练",
                "content": "机械能专题训练",
                "blockKind": "exam_title",
                "style": {"textAlign": "center", "fontWeight": "bold"},
            },
            {
                "id": "text-section",
                "title": "一、选择题",
                "content": "一、选择题",
                "blockKind": "section_title",
                "style": {"fontWeight": "bold"},
            },
        ],
        "nodes": [
            {"type": "text", "id": "node-title", "textBlockId": "text-title"},
            {"type": "text", "id": "node-section", "textBlockId": "text-section"},
            {"type": "knowledge", "id": "node-k", "knowledgeId": "k-1"},
            {"type": "question", "id": "node-q", "questionId": "q-1"},
        ],
        "headerFooter": {
            "headerEnabled": True,
            "headerText": "高二物理",
            "headerAlign": "center",
            "footerEnabled": True,
            "footerText": "Physics Vault",
            "footerAlign": "center",
            "showPageNumber": True,
        },
        "styleConfig": {
            "fontFamily": "songti",
            "fontSize": 12,
            "lineHeight": 1.55,
            "paragraphSpacing": 4,
            "questionSpacing": 10,
            "figureScale": 60,
            "pageMarginTop": 18,
            "pageMarginBottom": 18,
            "pageMarginLeft": 20,
            "pageMarginRight": 20,
            "pageSize": "A4",
            "pageOrientation": "portrait",
            "layoutMode": "paged-single",
            "optionLayout": "double",
        },
        "slideTemplate": "teach_practice_teach",
        "createdAt": "2026-08-03T00:00:00Z",
        "updatedAt": "2026-08-03T00:00:00Z",
    }


def _build_test_stack(tmp_path: Path) -> tuple[TestClient, LessonExportService, SQLiteImportTaskRepository, Path]:
    project_dir = tmp_path / "project"
    export_dir = project_dir / "data" / "exports"
    image_dir = project_dir / "data" / "assets" / "questions"
    image_dir.mkdir(parents=True)
    Image.new("RGB", (640, 360), "white").save(image_dir / "motion.png")

    repository = SQLiteImportTaskRepository(str(tmp_path / "tasks.sqlite3"))
    export_service = LessonExportService(
        repository,
        export_dir=export_dir,
        project_dir=project_dir,
        settings=LessonExportSettings(retention_days=7, max_snapshot_bytes=5 * 1024 * 1024),
    )
    queue_settings = TaskQueueSettings(enabled=False, max_retries=2)
    export_dispatcher = LessonExportDispatcher(export_service, queue_settings)
    pipeline = ImportPipelineService(
        task_repo=repository,
        pandoc=PandocAdapter(),
        cleaner=DocumentCleaningService(),
        parser=StructuredQuestionParsingService(),
    )
    app = FastAPI()
    app.include_router(build_lesson_exports_router(export_service, export_dispatcher))
    app.include_router(
        build_tasks_router(
            TaskCenterService(
                pipeline,
                lesson_export_service=export_service,
                export_dispatcher=export_dispatcher,
            )
        )
    )
    return TestClient(app), export_service, repository, project_dir


def test_word_export_snapshot_structure_and_safe_download(tmp_path: Path) -> None:
    from docx import Document

    client, _service, _repository, project_dir = _build_test_stack(tmp_path)
    response = client.post(
        "/api/exports/word",
        json={
            "lesson_package": _lesson_package("data/assets/questions/motion.png"),
            "include_answers": True,
            "include_analysis": True,
            "file_name": "../../CON.docx",
        },
    )

    assert response.status_code == 200
    task = response.json()
    assert task["status"] == "completed"
    assert task["progress"] == 100
    assert task["task_type"] == "word_export"
    events = client.get(f"/api/tasks/{task['task_id']}/events")
    assert events.status_code == 200
    assert [event["event_type"] for event in events.json()] == [
        "queued", "started", "started", "started", "completed",
    ]
    assert [event["phase"] for event in events.json()] == [
        "queued", "read", "export", "write", "complete",
    ]
    assert all(event["input_version"] == 1 for event in events.json())
    task_dir = project_dir / "data" / "exports" / task["task_id"]
    assert (task_dir / "snapshot.json").is_file()
    output_path = project_dir / task["result"]["result_file_path"]
    assert output_path.name == "physics-vault-export.docx"
    assert not list(task_dir.glob("*.tmp"))

    snapshot = json.loads((task_dir / "snapshot.json").read_text(encoding="utf-8"))
    assert snapshot["lesson_package"]["title"] == "机械能守恒练习"
    assert snapshot["options"] == {
        "include_answers": True,
        "include_analysis": True,
        "answer_position": "after_question",
    }

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "机械能守恒练习" not in text
    assert "学生版与教师版共用快照" not in text
    assert "机械能专题训练" in text
    assert "【答案】" in text
    assert "【详解】" in text
    assert "A. 动能为" in table_text
    assert document.styles["Normal"].font.name == "SimSun"
    assert document.styles["Normal"].font.size.pt == 10.5
    title_run = next(paragraph for paragraph in document.paragraphs if paragraph.text == "机械能专题训练").runs[0]
    section_run = next(paragraph for paragraph in document.paragraphs if paragraph.text == "一、选择题").runs[0]
    knowledge_run = next(paragraph for paragraph in document.paragraphs if paragraph.text == "机械能守恒").runs[0]
    knowledge_body_run = next(paragraph for paragraph in document.paragraphs if "只有重力做功" in paragraph.text).runs[0]
    answer_run = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("【答案】")).runs[0]
    analysis_run = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("【详解】")).runs[0]
    assert (title_run.font.name, title_run.font.size.pt, str(title_run.font.color.rgb)) == ("SimSun", 16, "000000")
    assert (section_run.font.name, section_run.font.size.pt, str(section_run.font.color.rgb)) == ("SimSun", 14, "000000")
    assert (knowledge_run.font.name, knowledge_run.font.size.pt, str(knowledge_run.font.color.rgb)) == ("SimSun", 14, "000000")
    assert (knowledge_body_run.font.name, knowledge_body_run.font.size.pt, str(knowledge_body_run.font.color.rgb)) == ("SimSun", 10.5, "000000")
    assert (answer_run.font.name, answer_run.font.size.pt, str(answer_run.font.color.rgb)) == ("KaiTi", 10.5, "000000")
    assert (analysis_run.font.name, analysis_run.font.size.pt, str(analysis_run.font.color.rgb)) == ("KaiTi", 10.5, "000000")
    question_paragraph = next(paragraph for paragraph in document.paragraphs if "质量为" in paragraph.text)
    assert all(run.bold in {None, False} for run in question_paragraph.runs)
    with zipfile.ZipFile(output_path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert document_xml.count("<m:oMath>") >= 5
    assert "<m:sSup>" in document_xml
    assert "<m:sSub>" in document_xml
    assert "<m:f>" in document_xml
    assert "<m:t>F</m:t>" in document_xml
    assert "<m:t>v</m:t>" in document_xml
    assert "<m:t>mgh</m:t>" in document_xml
    assert '<w:jc w:val="left"/>' in document_xml
    assert len(document.inline_shapes) == 1
    assert document.sections[0].header.paragraphs[0].text == "高二物理"
    assert "Physics Vault" in document.sections[0].footer.paragraphs[0].text

    download = client.get(f"/api/tasks/{task['task_id']}/download")
    assert download.status_code == 200
    assert download.content[:2] == b"PK"
    assert "physics-vault-export.docx" in download.headers["content-disposition"]


def test_word_export_can_place_answers_at_document_end(tmp_path: Path) -> None:
    from docx import Document

    client, _service, _repository, project_dir = _build_test_stack(tmp_path)
    response = client.post(
        "/api/exports/word",
        json={
            "lesson_package": _lesson_package("data/assets/questions/motion.png"),
            "include_answers": True,
            "include_analysis": True,
            "answer_position": "end",
        },
    )

    assert response.status_code == 200
    task = response.json()
    output_path = project_dir / task["result"]["result_file_path"]
    document = Document(output_path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    heading_index = paragraphs.index("参考答案与解析")
    answer_index = next(index for index, text in enumerate(paragraphs) if text.startswith("1. 【答案】"))
    assert heading_index < answer_index
    assert all("【答案】" not in text for text in paragraphs[:heading_index])


def test_pptx_export_preserves_template_teacher_content_and_images(tmp_path: Path) -> None:
    from pptx import Presentation

    client, service, _repository, project_dir = _build_test_stack(tmp_path)
    response = client.post(
        "/api/exports/pptx",
        json={
            "lesson_package": _lesson_package("data/assets/questions/motion.png"),
            "include_answers": True,
            "include_analysis": True,
        },
    )

    assert response.status_code == 200
    task = response.json()
    output_path = project_dir / task["result"]["result_file_path"]
    presentation = Presentation(output_path)
    assert len(presentation.slides) == 6  # cover, two title blocks, knowledge, question, teacher answer
    all_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "机械能守恒练习" in all_text
    assert "(1)/(2)mv²=mgh" in all_text
    assert "参考答案" in all_text
    assert "思路解析" in all_text
    assert any(shape.shape_type == 13 for slide in presentation.slides for shape in slide.shapes)

    # A redelivered message is idempotent: it does not create a second artifact.
    reexecuted = service.execute_export_task(task["task_id"])
    assert reexecuted.status == "completed"
    assert len(list(output_path.parent.glob("*.pptx"))) == 1


def test_download_rejects_result_path_outside_task_directory(tmp_path: Path) -> None:
    client, _service, repository, _project_dir = _build_test_stack(tmp_path)
    task = repository.create("word_export", {"filename": "unsafe.docx"})
    repository.mark_running(task.task_id)
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"PK unsafe")
    repository.mark_completed(
        task.task_id,
        {"result_file_path": str(outside)},
        result_file_path=str(outside),
    )

    response = client.get(f"/api/tasks/{task.task_id}/download")

    assert response.status_code == 404
    assert outside.is_file()


def test_cancelled_export_does_not_create_an_artifact(tmp_path: Path) -> None:
    _client, service, repository, project_dir = _build_test_stack(tmp_path)
    task = service.create_export_task(
        "word",
        LessonExportRequest(
            lesson_package=_lesson_package("data/assets/questions/motion.png"),
            include_answers=False,
            include_analysis=False,
        ),
    )
    repository.request_cancel(task.task_id)

    cancelled = service.execute_export_task(task.task_id)

    assert cancelled.status == "cancelled"
    assert not list((project_dir / "data" / "exports" / task.task_id).glob("*.docx"))


def test_failed_export_can_be_retried_from_task_center(tmp_path: Path) -> None:
    client, service, repository, _project_dir = _build_test_stack(tmp_path)
    original = service.create_export_task(
        "word",
        LessonExportRequest(lesson_package=_lesson_package("data/assets/questions/motion.png")),
    )
    repository.mark_failed(original.task_id, "simulated failure")

    response = client.post(f"/api/tasks/{original.task_id}/retry")

    assert response.status_code == 200
    payload = response.json()
    assert payload["original_task_id"] == original.task_id
    assert payload["task"]["task_id"] != original.task_id
    assert payload["task"]["status"] == "completed"
    assert payload["task"]["result_available"] is True


def test_cleanup_removes_only_expired_orphan_directories(tmp_path: Path) -> None:
    _client, service, _repository, project_dir = _build_test_stack(tmp_path)
    task = service.create_export_task(
        "word",
        LessonExportRequest(lesson_package=_lesson_package("data/assets/questions/motion.png")),
    )
    referenced = project_dir / "data" / "exports" / task.task_id
    orphan = project_dir / "data" / "exports" / "old-orphan"
    orphan.mkdir()
    old_timestamp = (datetime.now(UTC) - timedelta(days=20)).timestamp()
    os.utime(referenced, (old_timestamp, old_timestamp))
    os.utime(orphan, (old_timestamp, old_timestamp))

    removed = service.cleanup_expired_orphans()

    assert orphan.resolve() in removed
    assert not orphan.exists()
    assert referenced.exists()


def test_export_actor_runs_with_stub_broker(tmp_path: Path) -> None:
    _client, service, _repository, _project_dir = _build_test_stack(tmp_path)
    settings = TaskQueueSettings(enabled=True, max_retries=1, min_backoff_ms=0, max_backoff_ms=0)
    task = service.create_export_task(
        "word",
        LessonExportRequest(lesson_package=_lesson_package("data/assets/questions/motion.png")),
        max_attempts=settings.max_attempts,
    )
    broker = StubBroker(fail_fast_default=False)
    actors = build_lesson_export_actors(
        broker,
        settings,
        worker_factory=lambda: SimpleNamespace(lesson_export_service=service),
    )
    worker = Worker(broker, worker_threads=1, worker_timeout=20)
    worker.start()
    try:
        actors.run_lesson_export.send(task.task_id)
        broker.join(settings.queue_name, timeout=10_000, fail_fast=False)
        broker.join(settings.control_queue_name, timeout=10_000, fail_fast=False)
    finally:
        worker.stop()
        worker.join()

    completed = service.get_task(task.task_id)
    assert completed.status == "completed"
    assert completed.result_file_path
