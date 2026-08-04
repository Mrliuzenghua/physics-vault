from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import shutil
from dataclasses import asdict
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from ..config import LessonExportSettings
from ..paths import default_exports_dir, project_root
from ..repositories.import_tasks import ImportTask
from ..schemas.lesson_exports import LessonExportRequest


ExportFormat = Literal["word", "pptx"]

_EXPORT_TASK_TYPES = {"word": "word_export", "pptx": "pptx_export"}
_EXPORT_SUFFIXES = {"word": ".docx", "pptx": ".pptx"}
_EXPORT_MIME_TYPES = {
    "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_LESSON_EXPORT_SNAPSHOT_SCHEMA = "physics-vault/lesson-export-snapshot/v3"
_WINDOWS_RESERVED_NAMES = {
    "CON",
    "PRN",
    "AUX",
    "NUL",
    *(f"COM{index}" for index in range(1, 10)),
    *(f"LPT{index}" for index in range(1, 10)),
}
_FIGURE_RE = re.compile(r"!\[fig:([^\]]+)\]")
_FORMULA_RE = re.compile(r"\$\$([\s\S]+?)\$\$|\$([^$]+?)\$|\\\[([\s\S]+?)\\\]|\\\(([^)]+?)\\\)")


class ExportCancelled(RuntimeError):
    pass


class LessonExportService:
    """Create immutable lesson snapshots and deterministic Office artifacts."""

    def __init__(
        self,
        task_repository: Any,
        *,
        export_dir: Path | None = None,
        project_dir: Path | None = None,
        settings: LessonExportSettings | None = None,
    ) -> None:
        self._repository = task_repository
        self._export_dir = (export_dir or default_exports_dir()).resolve()
        self._project_dir = (project_dir or project_root()).resolve()
        self.settings = settings or LessonExportSettings.from_env()
        self._export_dir.mkdir(parents=True, exist_ok=True)

    def create_export_task(
        self,
        export_format: ExportFormat,
        payload: LessonExportRequest,
        *,
        max_attempts: int = 1,
        request_context: dict[str, Any] | None = None,
        force_new: bool = False,
    ) -> ImportTask:
        snapshot = {
            "schema": _LESSON_EXPORT_SNAPSHOT_SCHEMA,
            "export_format": export_format,
            "lesson_package": payload.lesson_package,
            "options": {
                "include_answers": payload.include_answers,
                "include_analysis": payload.include_analysis,
                "answer_position": payload.answer_position,
            },
            "requested_file_name": payload.file_name,
            "created_at": datetime.now(UTC).isoformat(),
        }
        snapshot_bytes = json.dumps(snapshot, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        if len(snapshot_bytes) > self.settings.max_snapshot_bytes:
            raise ValueError(
                f"Export snapshot is too large ({len(snapshot_bytes)} bytes; "
                f"limit {self.settings.max_snapshot_bytes} bytes)"
            )

        lesson = payload.lesson_package
        title = str(lesson.get("title") or "未命名物理学案")
        input_summary = {
            "lesson_id": str(lesson.get("id") or ""),
            "revision": int(lesson.get("revision") or 0),
            "title": title,
            "filename": _safe_file_stem(payload.file_name or title, "physics-vault-export")
            + _EXPORT_SUFFIXES[export_format],
            "export_format": export_format,
            "include_answers": payload.include_answers,
            "include_analysis": payload.include_analysis,
            "answer_position": payload.answer_position,
            "question_count": len(lesson.get("questions") or []),
        }
        if request_context:
            input_summary["request_context"] = {
                key: str(value)
                for key, value in request_context.items()
                if value is not None and key in {"source", "session_id", "operator"}
            }
        idempotency_payload = {
            "schema": snapshot["schema"],
            "export_format": export_format,
            "lesson_package": payload.lesson_package,
            "options": snapshot["options"],
            "requested_file_name": payload.file_name,
        }
        idempotency_key = "lesson-export:" + hashlib.sha256(
            json.dumps(
                idempotency_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            ).encode("utf-8")
        ).hexdigest()
        create_or_get = getattr(self._repository, "create_or_get", None)
        created = True
        if callable(create_or_get) and not force_new:
            task, created = create_or_get(
                _EXPORT_TASK_TYPES[export_format],
                input_summary,
                max_attempts=max(1, int(max_attempts)),
                idempotency_key=idempotency_key,
            )
            if not created and self._export_task_is_reusable(task):
                return task
        if force_new or not created:
            task = self._repository.create(
                _EXPORT_TASK_TYPES[export_format],
                input_summary,
                max_attempts=max(1, int(max_attempts)),
                idempotency_key=(
                    f"{idempotency_key}:retry:{uuid4().hex}"
                    if force_new
                    else idempotency_key
                ),
            )
        elif not callable(create_or_get):
            task = self._repository.create(
                _EXPORT_TASK_TYPES[export_format],
                input_summary,
                max_attempts=max(1, int(max_attempts)),
                idempotency_key=idempotency_key,
            )
        try:
            task_dir = self._task_dir(task.task_id)
            task_dir.mkdir(parents=True, exist_ok=False)
            _atomic_write_bytes(task_dir / "snapshot.json", snapshot_bytes)
        except Exception as exc:
            self._repository.mark_failed(
                task.task_id,
                str(exc),
                error_type=type(exc).__name__,
                user_message="无法保存导出快照，请检查导出目录后重试。",
                technical_details=str(exc),
                retryable=False,
            )
            raise
        self.cleanup_expired_orphans()
        return self._require_task(task.task_id)

    def clone_for_retry(self, task_id: str, *, max_attempts: int = 1) -> ImportTask:
        original = self._require_export_task(task_id)
        if original.status not in {"failed", "cancelled"}:
            raise ValueError("Only failed or cancelled export tasks can be retried")
        snapshot = self._read_snapshot(task_id)
        return self.create_export_task(
            str(snapshot["export_format"]),  # type: ignore[arg-type]
            LessonExportRequest(
                lesson_package=dict(snapshot["lesson_package"]),
                include_answers=bool(snapshot.get("options", {}).get("include_answers")),
                include_analysis=bool(snapshot.get("options", {}).get("include_analysis")),
                answer_position=str(snapshot.get("options", {}).get("answer_position") or "after_question"),  # type: ignore[arg-type]
                file_name=snapshot.get("requested_file_name"),
            ),
            max_attempts=max_attempts,
            force_new=True,
        )

    def _export_task_is_reusable(self, task: ImportTask) -> bool:
        if task.status in {"pending", "running", "retrying", "cancel_requested"}:
            return self._task_dir(task.task_id).joinpath("snapshot.json").is_file()
        if task.status != "completed":
            return False
        try:
            self.resolve_result_file(task.task_id)
        except (FileNotFoundError, ValueError):
            return False
        return True

    def execute_export_task(self, task_id: str) -> ImportTask:
        current = self._require_export_task(task_id)
        if current.status in {"completed", "failed", "cancelled"}:
            return current
        if current.status == "cancel_requested":
            return self._repository.mark_cancelled(task_id, "已取消")
        claimed = self._repository.claim(task_id, current_step="读取不可变快照")
        if claimed is None:
            return self._require_task(task_id)

        snapshot = self._read_snapshot(task_id)
        export_format = str(snapshot.get("export_format") or "")
        if export_format not in _EXPORT_TASK_TYPES:
            raise ValueError(f"Unsupported export format in snapshot: {export_format}")
        lesson = snapshot.get("lesson_package")
        if not isinstance(lesson, dict):
            raise ValueError("Export snapshot does not contain a lesson package")
        options = snapshot.get("options") if isinstance(snapshot.get("options"), dict) else {}
        requested_name = snapshot.get("requested_file_name")
        title = str(lesson.get("title") or "未命名物理学案")
        final_name = _safe_file_stem(str(requested_name or title), "physics-vault-export") + _EXPORT_SUFFIXES[export_format]
        task_dir = self._task_dir(task_id)
        final_path = task_dir / final_name
        temporary_path = task_dir / f".{final_name}.{uuid4().hex}.tmp"

        try:
            self._check_cancelled(task_id)
            self._repository.update_progress(task_id, 15, "准备版式与媒体")
            if export_format == "word":
                warnings = self._generate_docx(
                    lesson,
                    options,
                    temporary_path,
                    lambda: self._check_cancelled(task_id),
                )
            else:
                warnings = self._generate_pptx(
                    lesson,
                    options,
                    temporary_path,
                    lambda: self._check_cancelled(task_id),
                )
            self._repository.update_progress(task_id, 90, "校验并保存产物")
            _verify_office_package(temporary_path, export_format)
            self._check_cancelled(task_id)
            os.replace(temporary_path, final_path)
            relative_path = _portable_path(final_path, self._project_dir)
            result = {
                "task_id": task_id,
                "status": "completed",
                "export_format": export_format,
                "filename": final_name,
                "mime_type": _EXPORT_MIME_TYPES[export_format],
                "size": final_path.stat().st_size,
                "sha256": _sha256_file(final_path),
                "result_file_path": relative_path,
                "snapshot_path": _portable_path(task_dir / "snapshot.json", self._project_dir),
                "download_url": f"/api/tasks/{task_id}/download",
                "warnings": warnings,
            }
            return self._repository.mark_completed(
                task_id,
                result,
                result_file_path=relative_path,
            )
        except ExportCancelled:
            temporary_path.unlink(missing_ok=True)
            return self._repository.mark_cancelled(task_id, "已取消")
        except Exception:
            temporary_path.unlink(missing_ok=True)
            raise

    def mark_retrying(self, task_id: str, exc: BaseException) -> ImportTask:
        return self._repository.mark_retrying(
            task_id,
            str(exc),
            error_type=type(exc).__name__,
            user_message="导出暂时失败，正在自动重试。",
            technical_details=str(exc),
        )

    def mark_failed(self, task_id: str, exc: BaseException, *, retryable: bool = False) -> ImportTask:
        current = self._require_task(task_id)
        if current.status in {"completed", "failed", "cancelled"}:
            return current
        return self._repository.mark_failed(
            task_id,
            str(exc),
            error_type=type(exc).__name__,
            user_message="服务端导出失败。你可以在任务中心重试，或回到组卷页使用浏览器导出。",
            technical_details=str(exc),
            retryable=retryable,
        )

    def get_task(self, task_id: str) -> ImportTask:
        return self._require_task(task_id)

    def set_message_id(self, task_id: str, message_id: str) -> ImportTask:
        setter = getattr(self._repository, "set_message_id", None)
        return setter(task_id, message_id) if callable(setter) else self._require_task(task_id)

    def resolve_result_file(self, task_id: str) -> Path:
        task = self._require_export_task(task_id)
        if task.status != "completed":
            raise FileNotFoundError("Export task has not completed")
        raw = task.result_file_path or str((task.result or {}).get("result_file_path") or "")
        if not raw:
            raise FileNotFoundError("Export task has no result file")
        candidate = Path(raw)
        resolved = (self._project_dir / candidate).resolve() if not candidate.is_absolute() else candidate.resolve()
        task_dir = self._task_dir(task_id)
        if task_dir not in resolved.parents or resolved.suffix.lower() not in {".docx", ".pptx"}:
            raise FileNotFoundError("Unsafe export result path")
        if not resolved.is_file():
            raise FileNotFoundError("Export result has expired")
        return resolved

    def cleanup_expired_orphans(self, now: datetime | None = None) -> list[Path]:
        """Remove only old, unreferenced task directories; live task records always win."""
        cutoff = (now or datetime.now(UTC)) - timedelta(days=self.settings.retention_days)
        removed: list[Path] = []
        for candidate in self._export_dir.iterdir():
            if not candidate.is_dir():
                continue
            try:
                modified = datetime.fromtimestamp(candidate.stat().st_mtime, UTC)
                if modified >= cutoff or self._repository.get(candidate.name) is not None:
                    continue
                resolved = candidate.resolve()
                if resolved.parent != self._export_dir:
                    continue
                shutil.rmtree(resolved)
                removed.append(resolved)
            except OSError:
                continue
        return removed

    def _generate_docx(
        self,
        lesson: dict[str, Any],
        options: dict[str, Any],
        output_path: Path,
        check_cancelled: Any,
    ) -> list[str]:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor

        warnings: list[str] = []
        document = Document()
        section = document.sections[0]
        style = lesson.get("styleConfig") if isinstance(lesson.get("styleConfig"), dict) else {}
        page_size = str(style.get("pageSize") or "A4")
        width_mm, height_mm = ((297, 420) if page_size == "A3" else (210, 297))
        if str(style.get("pageOrientation") or "portrait") == "landscape":
            width_mm, height_mm = height_mm, width_mm
            section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(width_mm)
        section.page_height = Mm(height_mm)
        section.top_margin = Mm(float(style.get("pageMarginTop") or 18))
        section.bottom_margin = Mm(float(style.get("pageMarginBottom") or 18))
        section.left_margin = Mm(float(style.get("pageMarginLeft") or 20))
        section.right_margin = Mm(float(style.get("pageMarginRight") or 20))

        # Server Word export follows the standard printable exam style:
        # SimSun body text at 5 hao (10.5 pt), regardless of screen preview scale.
        font_name = "SimSun"
        answer_font = "KaiTi"
        body_size = 10.5
        title_size = 16
        small_title_size = 14
        word_color = "000000"
        answer_position = str(options.get("answer_position") or "after_question")
        trailing_answer_blocks: list[tuple[int, dict[str, Any]]] = []
        normal = document.styles["Normal"]
        normal.font.name = font_name
        normal.font.size = Pt(body_size)
        normal.font.color.rgb = RGBColor.from_string(word_color)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        normal.paragraph_format.line_spacing = float(style.get("lineHeight") or 1.55)
        normal.paragraph_format.space_after = Pt(float(style.get("paragraphSpacing") or 4))

        header_footer = lesson.get("headerFooter") if isinstance(lesson.get("headerFooter"), dict) else {}
        if header_footer.get("headerEnabled"):
            paragraph = section.header.paragraphs[0]
            paragraph.alignment = _docx_alignment(str(header_footer.get("headerAlign") or "center"))
            _add_docx_text(paragraph, str(header_footer.get("headerText") or ""), font_name, 9, color=word_color)
        footer = section.footer.paragraphs[0]
        footer.alignment = _docx_alignment(str(header_footer.get("footerAlign") or "center"))
        if header_footer.get("footerEnabled"):
            _add_docx_text(footer, str(header_footer.get("footerText") or ""), font_name, 9, color=word_color)
        if header_footer.get("showPageNumber"):
            if footer.text:
                _format_run(footer.add_run("  ·  "), font_name, 9, color=word_color)
            _format_run(footer.add_run("第 "), font_name, 9, color=word_color)
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
            _format_run(footer.add_run(" 页"), font_name, 9, color=word_color)

        question_map = {
            str(item.get("question_id")): item
            for item in lesson.get("questions") or []
            if isinstance(item, dict) and item.get("question_id")
        }
        knowledge_map = {
            str(item.get("id")): item
            for item in lesson.get("knowledgeCards") or []
            if isinstance(item, dict) and item.get("id")
        }
        text_map = {
            str(item.get("id")): item
            for item in lesson.get("textBlocks") or []
            if isinstance(item, dict) and item.get("id")
        }
        question_number = 0
        for node in lesson.get("nodes") or []:
            check_cancelled()
            if not isinstance(node, dict):
                continue
            node_type = str(node.get("type") or "")
            if node_type == "page_break":
                document.add_page_break()
                continue
            if node_type == "knowledge":
                card = knowledge_map.get(str(node.get("knowledgeId") or ""))
                if not card:
                    warnings.append(f"Missing knowledge card for node {node.get('id')}")
                    continue
                heading = document.add_paragraph()
                heading.paragraph_format.space_before = Pt(10)
                heading.paragraph_format.space_after = Pt(4)
                heading.paragraph_format.keep_with_next = True
                _add_docx_text(heading, str(card.get("title") or "知识梳理"), font_name, small_title_size, bold=True, color=word_color)
                _add_docx_text(document.add_paragraph(), str(card.get("summary") or ""), font_name, body_size, color=word_color)
                for point in card.get("points") or []:
                    _add_docx_text(document.add_paragraph(style="List Bullet"), str(point), font_name, body_size, color=word_color)
                continue
            if node_type == "text":
                block = text_map.get(str(node.get("textBlockId") or ""))
                if not block:
                    warnings.append(f"Missing text block for node {node.get('id')}")
                    continue
                block_kind = str(block.get("blockKind") or "body")
                paragraph = document.add_paragraph()
                block_style = block.get("style") if isinstance(block.get("style"), dict) else {}
                if block_style.get("textAlign"):
                    paragraph.alignment = _docx_alignment(str(block_style["textAlign"]))
                elif block_kind == "exam_title":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                is_heading = block_kind in {"exam_title", "section_title"}
                paragraph.paragraph_format.keep_with_next = is_heading
                paragraph.paragraph_format.space_before = Pt(10 if is_heading else 4)
                paragraph.paragraph_format.space_after = Pt(6 if is_heading else 4)
                block_size = title_size if block_kind == "exam_title" else small_title_size if block_kind == "section_title" else body_size
                _add_docx_text(
                    paragraph,
                    str(block.get("content") or block.get("title") or ""),
                    font_name,
                    block_size,
                    bold=is_heading or block_style.get("fontWeight") == "bold",
                    color=word_color,
                )
                continue
            if node_type != "question":
                continue
            question = question_map.get(str(node.get("questionId") or ""))
            if not question:
                warnings.append(f"Missing question for node {node.get('id')}")
                continue
            question_number += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(float(style.get("questionSpacing") or 10))
            _format_run(paragraph.add_run(f"{question_number}. "), font_name, body_size)
            _add_docx_text(paragraph, str(question.get("title") or question.get("stem_text") or ""), font_name, body_size)

            figures = [item for item in question.get("figures") or [] if isinstance(item, dict)]
            for figure in figures:
                image = self._read_image(str(figure.get("local_path") or ""))
                if image is None:
                    warnings.append(f"Image unavailable: {figure.get('local_path') or figure.get('fig_uuid')}")
                    continue
                image_stream, image_width, image_height = image
                scale = _bounded_number(figure.get("display_scale") or style.get("figureScale"), 25, 100, 60)
                available_inches = max(1.0, (section.page_width - section.left_margin - section.right_margin) / 914400)
                width_inches = min(available_inches * scale / 100, image_width / 96)
                picture = document.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.LEFT
                picture.add_run().add_picture(image_stream, width=int(width_inches * 914400))
                if figure.get("caption"):
                    caption = document.add_paragraph()
                    caption.alignment = picture.alignment
                    _add_docx_text(caption, str(figure["caption"]), font_name, max(9, body_size - 2), italic=True, color=word_color)

            options_list = [item for item in question.get("options") or [] if isinstance(item, dict)]
            option_layout = str(style.get("optionLayout") or "auto")
            use_two_columns = option_layout == "double" or (
                option_layout == "auto"
                and len(options_list) in {4, 6}
                and all(len(_office_text(str(item.get("content") or ""))) <= 32 for item in options_list)
            )
            if use_two_columns and options_list:
                table = document.add_table(rows=(len(options_list) + 1) // 2, cols=2)
                table.autofit = False
                for index, option in enumerate(options_list):
                    cell = table.cell(index // 2, index % 2)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    label = str(option.get("opt") or chr(65 + index))
                    _format_run(cell.paragraphs[0].add_run(f"{label}. "), font_name, body_size)
                    _add_docx_text(cell.paragraphs[0], str(option.get("content") or ""), font_name, body_size)
                borders = OxmlElement("w:tblBorders")
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    element = OxmlElement(f"w:{edge}")
                    element.set(qn("w:val"), "nil")
                    borders.append(element)
                table._tbl.tblPr.append(borders)
            else:
                for index, option in enumerate(options_list):
                    label = str(option.get("opt") or chr(65 + index))
                    option_paragraph = document.add_paragraph()
                    _format_run(option_paragraph.add_run(f"{label}. "), font_name, body_size)
                    _add_docx_text(option_paragraph, str(option.get("content") or ""), font_name, body_size)
            if answer_position == "end":
                trailing_answer_blocks.append((question_number, question))
            else:
                _add_docx_answer_block(document, question, None, options, answer_font, body_size)

        if trailing_answer_blocks and (options.get("include_answers") or options.get("include_analysis")):
            document.add_page_break()
            heading = document.add_paragraph()
            heading.paragraph_format.space_after = Pt(8)
            heading_text = "参考答案与解析" if options.get("include_analysis") else "参考答案"
            _format_run(heading.add_run(heading_text), font_name, small_title_size, bold=True, color=word_color)
            for item_number, question in trailing_answer_blocks:
                _add_docx_answer_block(document, question, item_number, options, answer_font, body_size)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
        return warnings

    def _generate_pptx(
        self,
        lesson: dict[str, Any],
        options: dict[str, Any],
        output_path: Path,
        check_cancelled: Any,
    ) -> list[str]:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        warnings: list[str] = []
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]

        def add_text(slide: Any, text: str, x: float, y: float, w: float, h: float, *, size: int, bold: bool = False, color: str = "10233F", align: Any = PP_ALIGN.LEFT) -> Any:
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            frame = box.text_frame
            frame.clear()
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = _office_text(text)
            paragraph.alignment = align
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            return box

        def add_heading(slide: Any, title: str, eyebrow: str, page: int) -> None:
            add_text(slide, eyebrow, 0.72, 0.34, 4.8, 0.3, size=12, bold=True, color="2567B8")
            add_text(slide, title, 0.72, 0.78, 11.9, 0.68, size=36, bold=True)
            add_text(slide, str(page), 12.1, 7.1, 0.45, 0.18, size=8, color="53647A", align=PP_ALIGN.RIGHT)

        cover = presentation.slides.add_slide(blank)
        add_text(cover, "PHYSICS VAULT · CLASSROOM KIT", 0.9, 1.1, 7.5, 0.35, size=16, bold=True, color="2567B8")
        add_text(cover, str(lesson.get("title") or "未命名物理课件"), 0.9, 1.75, 11.3, 1.25, size=54, bold=True)
        add_text(cover, str(lesson.get("subtitle") or ""), 0.94, 3.3, 10.8, 0.7, size=24, color="53647A")
        add_text(cover, f"{len(lesson.get('questions') or [])} 道题目", 0.94, 6.25, 4.0, 0.35, size=18, color="53647A")

        question_map = {
            str(item.get("question_id")): item
            for item in lesson.get("questions") or []
            if isinstance(item, dict) and item.get("question_id")
        }
        knowledge_map = {
            str(item.get("id")): item
            for item in lesson.get("knowledgeCards") or []
            if isinstance(item, dict) and item.get("id")
        }
        text_map = {
            str(item.get("id")): item
            for item in lesson.get("textBlocks") or []
            if isinstance(item, dict) and item.get("id")
        }
        nodes = [node for node in lesson.get("nodes") or [] if isinstance(node, dict)]
        template = str(lesson.get("slideTemplate") or "teach_practice_teach")
        if template == "practice_only":
            nodes = [node for node in nodes if node.get("type") == "question"]
        elif template == "teach_then_practice":
            nodes = [node for node in nodes if node.get("type") in {"knowledge", "text"}] + [
                node for node in nodes if node.get("type") == "question"
            ]

        page = 1
        question_number = 0
        for node in nodes:
            check_cancelled()
            node_type = str(node.get("type") or "")
            if node_type in {"page_break", "separator"}:
                continue
            if node_type == "knowledge":
                card = knowledge_map.get(str(node.get("knowledgeId") or ""))
                if not card:
                    warnings.append(f"Missing knowledge card for node {node.get('id')}")
                    continue
                lines = [str(card.get("summary") or "")] + [f"• {item}" for item in card.get("points") or []]
                for chunk_index, chunk in enumerate(_chunk_lines(lines, 650)):
                    page += 1
                    slide = presentation.slides.add_slide(blank)
                    add_heading(slide, str(card.get("title") or "知识梳理") + ("（续）" if chunk_index else ""), "KNOWLEDGE MAP", page)
                    add_text(slide, "\n\n".join(chunk), 0.9, 1.75, 11.5, 4.9, size=22, color="10233F")
                continue
            if node_type == "text":
                block = text_map.get(str(node.get("textBlockId") or ""))
                if not block:
                    warnings.append(f"Missing text block for node {node.get('id')}")
                    continue
                chunks = _chunk_lines(str(block.get("content") or "").splitlines() or [str(block.get("title") or "")], 650)
                for chunk_index, chunk in enumerate(chunks):
                    page += 1
                    slide = presentation.slides.add_slide(blank)
                    add_heading(slide, str(block.get("title") or "教学说明") + ("（续）" if chunk_index else ""), "TEACHING FLOW", page)
                    add_text(slide, "\n\n".join(chunk), 0.9, 1.8, 11.5, 4.8, size=22)
                continue
            if node_type != "question":
                continue
            question = question_map.get(str(node.get("questionId") or ""))
            if not question:
                warnings.append(f"Missing question for node {node.get('id')}")
                continue
            question_number += 1
            question_lines = [str(question.get("title") or question.get("stem_text") or "")]
            question_lines.extend(
                f"{item.get('opt') or chr(65 + index)}. {item.get('content') or ''}"
                for index, item in enumerate(question.get("options") or [])
                if isinstance(item, dict)
            )
            chunks = _chunk_lines(question_lines, 540)
            for chunk_index, chunk in enumerate(chunks):
                page += 1
                slide = presentation.slides.add_slide(blank)
                title = f"题目 {question_number}" + ("（续）" if chunk_index else "")
                add_heading(slide, title, _question_knowledge(question), page)
                image = None
                if chunk_index == 0:
                    figures = [item for item in question.get("figures") or [] if isinstance(item, dict)]
                    if figures:
                        image = self._read_image(str(figures[0].get("local_path") or ""))
                        if image is None:
                            warnings.append(f"Image unavailable: {figures[0].get('local_path')}")
                text_width = 7.4 if image else 11.5
                add_text(slide, "\n\n".join(chunk), 0.9, 1.75, text_width, 4.95, size=20, bold=chunk_index == 0)
                if image:
                    stream, width_px, height_px = image
                    max_w, max_h = 3.55, 3.8
                    scale = min(max_w / width_px, max_h / height_px)
                    width = width_px * scale
                    height = height_px * scale
                    slide.shapes.add_picture(stream, Inches(8.8), Inches(2.0), Inches(width), Inches(height))
            answer_parts: list[str] = []
            if options.get("include_answers") and question.get("answer"):
                answer_parts.append(f"参考答案\n{question['answer']}")
            if options.get("include_analysis") and question.get("analysis"):
                answer_parts.append(f"思路解析\n{question['analysis']}")
            answer_chunks = _chunk_lines(answer_parts, 650) if answer_parts else []
            for chunk_index, chunk in enumerate(answer_chunks):
                page += 1
                slide = presentation.slides.add_slide(blank)
                add_heading(slide, f"题目 {question_number} · 讲解" + ("（续）" if chunk_index else ""), _question_knowledge(question), page)
                add_text(slide, "\n\n".join(chunk), 0.9, 1.75, 11.5, 4.9, size=22, color="0E7A58")

        presentation.core_properties.author = "Physics Vault"
        presentation.core_properties.title = str(lesson.get("title") or "")
        presentation.core_properties.subject = str(lesson.get("subtitle") or "")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(output_path)
        return warnings

    def _read_image(self, raw_path: str) -> tuple[io.BytesIO, int, int] | None:
        from PIL import Image, ImageOps

        data: bytes
        if raw_path.startswith("data:image/"):
            try:
                data = base64.b64decode(raw_path.split(",", 1)[1], validate=True)
            except (IndexError, ValueError):
                return None
        else:
            path = _resolve_project_image(raw_path, self._project_dir)
            if path is None:
                return None
            try:
                data = path.read_bytes()
            except OSError:
                return None
        if len(data) > 25 * 1024 * 1024:
            return None
        try:
            with Image.open(io.BytesIO(data)) as source:
                image = ImageOps.exif_transpose(source)
                width, height = image.size
                if image.mode not in {"RGB", "RGBA"}:
                    image = image.convert("RGBA" if "transparency" in image.info else "RGB")
                stream = io.BytesIO()
                image.save(stream, format="PNG")
                stream.seek(0)
                return stream, max(1, width), max(1, height)
        except Exception:
            return None

    def _check_cancelled(self, task_id: str) -> None:
        current = self._require_task(task_id)
        if current.status in {"cancel_requested", "cancelled"}:
            raise ExportCancelled("Export was cancelled")

    def _read_snapshot(self, task_id: str) -> dict[str, Any]:
        snapshot_path = self._task_dir(task_id) / "snapshot.json"
        if not snapshot_path.is_file():
            raise FileNotFoundError("Immutable export snapshot is missing")
        with snapshot_path.open("r", encoding="utf-8") as handle:
            value = json.load(handle)
        if not isinstance(value, dict):
            raise ValueError("Invalid export snapshot")
        return value

    def _task_dir(self, task_id: str) -> Path:
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", task_id):
            raise ValueError("Invalid task id")
        candidate = (self._export_dir / task_id).resolve()
        if candidate.parent != self._export_dir:
            raise ValueError("Invalid export task path")
        return candidate

    def _require_task(self, task_id: str) -> ImportTask:
        task = self._repository.get(task_id)
        if task is None:
            raise KeyError(task_id)
        return task

    def _require_export_task(self, task_id: str) -> ImportTask:
        task = self._require_task(task_id)
        if task.task_type not in set(_EXPORT_TASK_TYPES.values()):
            raise ValueError("Task is not a lesson export")
        return task


def serialize_export_task(task: ImportTask) -> dict[str, Any]:
    return asdict(task)


def _atomic_write_bytes(path: Path, payload: bytes) -> None:
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        with temporary.open("xb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _safe_file_stem(value: str, fallback: str) -> str:
    stem = Path(str(value or "").replace("\\", "/")).name
    stem = re.sub(r"\.(docx|pptx)$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"[\x00-\x1f<>:\"/\\|?*]", "-", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" .")[:96]
    if not stem or stem.upper() in _WINDOWS_RESERVED_NAMES:
        return fallback
    return stem


def _resolve_project_image(raw_path: str, root: Path) -> Path | None:
    value = str(raw_path or "").strip()
    if not value or re.match(r"^(https?:|blob:)", value, flags=re.IGNORECASE):
        return None
    value = value.replace("\\", "/").removeprefix("file:///").removeprefix("file://")
    match = re.search(r"(?:^|/)physics-vault/(.+)$", value, flags=re.IGNORECASE)
    if match:
        value = match.group(1)
    value = value.lstrip("/").removeprefix("./")
    if value.startswith("files/"):
        value = value[6:]
    if value.startswith("assets/questions/"):
        value = f"data/{value}"
    elif value.startswith("import-batches/"):
        value = f"data/{value}"
    elif "/" not in value and re.search(r"\.(png|jpe?g|gif|webp|bmp)$", value, flags=re.IGNORECASE):
        value = f"data/assets/questions/{value}"
    candidate = (root / value).resolve()
    if root not in candidate.parents or not candidate.is_file():
        return None
    return candidate


def _add_docx_answer_block(
    document: Any,
    question: dict[str, Any],
    question_number: int | None,
    options: dict[str, Any],
    font_name: str,
    body_size: float,
) -> None:
    from docx.shared import Pt

    if options.get("include_answers") and question.get("answer"):
        answer = document.add_paragraph()
        answer.paragraph_format.space_before = Pt(4)
        prefix = f"{question_number}. 答案：" if question_number is not None else "答案："
        _format_run(answer.add_run(prefix), font_name, body_size, color="000000")
        _add_docx_text(answer, str(question["answer"]), font_name, body_size, color="000000")
    if options.get("include_analysis") and question.get("analysis"):
        analysis = document.add_paragraph()
        prefix = f"{question_number}. 解析：" if question_number is not None and not question.get("answer") else "解析："
        _format_run(analysis.add_run(prefix), font_name, body_size, color="000000")
        _add_docx_text(analysis, str(question["analysis"]), font_name, body_size, color="000000")


def _office_text(value: str) -> str:
    text = _FIGURE_RE.sub("", str(value or ""))

    def replace_formula(match: re.Match[str]) -> str:
        formula = next((group for group in match.groups() if group is not None), "")
        return _linearize_latex(formula)

    text = _FORMULA_RE.sub(replace_formula, text)
    text = re.sub(r"^#{1,4}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*\*([^*\n]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`\n]+)`", r"\1", text)
    return re.sub(r"[ \t]+", " ", text).strip()


def _office_text_parts(value: str) -> list[tuple[str, str]]:
    text = _FIGURE_RE.sub("", str(value or ""))
    text = re.sub(r"^#{1,4}\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"\*\*([^*\n]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`\n]+)`", r"\1", text)
    text = re.sub(r"[ \t]+", " ", text).strip()
    if not text:
        return []

    parts: list[tuple[str, str]] = []
    cursor = 0
    for match in _FORMULA_RE.finditer(text):
        plain = text[cursor : match.start()]
        if plain:
            parts.append(("text", plain))
        formula = next((group for group in match.groups() if group is not None), "").strip()
        if formula:
            parts.append(("math", formula))
        cursor = match.end()
    tail = text[cursor:]
    if tail:
        parts.append(("text", tail))
    return parts


def _add_docx_text(
    paragraph: Any,
    value: str,
    font: str,
    size: float,
    *,
    bold: bool = False,
    italic: bool = False,
    color: str | None = None,
) -> None:
    for kind, text in _office_text_parts(value):
        if kind == "math":
            paragraph._p.append(_latex_to_omml(text))
            continue
        run = paragraph.add_run(text)
        _format_run(run, font, size, bold=bold, italic=italic, color=color)


def _latex_to_omml(value: str) -> Any:
    from docx.oxml import OxmlElement

    math = OxmlElement("m:oMath")
    children = _latex_to_math_elements(value)
    if not children:
        children = [_math_run("")]
    for child in children:
        math.append(child)
    return math


def _latex_to_math_elements(value: str) -> list[Any]:
    source = str(value or "").strip()
    index = 0

    def parse_sequence(stop: str | None = None) -> list[Any]:
        nonlocal index
        items: list[Any] = []
        while index < len(source):
            current = source[index]
            if stop is not None and current == stop:
                index += 1
                break
            if current == "}":
                break
            if current.isspace():
                while index < len(source) and source[index].isspace():
                    index += 1
                items.append(_math_run(" "))
                continue
            if current in {"_", "^"}:
                items.append(_math_run(current))
                index += 1
                continue

            base = parse_atom()
            subscript: list[Any] | None = None
            superscript: list[Any] | None = None
            while index < len(source) and source[index] in {"_", "^"}:
                marker = source[index]
                index += 1
                script = parse_script_atom()
                if marker == "_":
                    subscript = script
                else:
                    superscript = script
            if subscript is not None or superscript is not None:
                items.append(_math_script(base, subscript, superscript))
            else:
                items.extend(base)
        return items

    def parse_script_atom() -> list[Any]:
        nonlocal index
        while index < len(source) and source[index].isspace():
            index += 1
        if index >= len(source):
            return [_math_run("")]
        if source[index] == "{":
            index += 1
            return parse_sequence("}")
        if source[index] == "\\":
            return parse_command()
        char = source[index]
        index += 1
        return [_math_run(char)]

    def parse_required_group() -> list[Any]:
        nonlocal index
        while index < len(source) and source[index].isspace():
            index += 1
        if index < len(source) and source[index] == "{":
            index += 1
            return parse_sequence("}")
        return parse_script_atom()

    def parse_atom() -> list[Any]:
        nonlocal index
        if source[index] == "{":
            index += 1
            return parse_sequence("}")
        if source[index] == "\\":
            return parse_command()

        start = index
        while index < len(source) and source[index] not in "\\{}_^" and not source[index].isspace():
            index += 1
        return [_math_run(source[start:index])]

    def parse_command() -> list[Any]:
        nonlocal index
        index += 1
        start = index
        while index < len(source) and source[index].isalpha():
            index += 1
        command = source[start:index]
        if not command and index < len(source):
            command = source[index]
            index += 1

        if command in {"frac", "dfrac", "tfrac"}:
            return [_math_fraction(parse_required_group(), parse_required_group())]
        if command == "sqrt":
            return [_math_radical(parse_required_group())]
        if command in {"text", "mathrm", "mathbf", "operatorname"}:
            return parse_required_group()
        if command in {"left", "right"}:
            return parse_script_atom() if index < len(source) else []
        if command in {"quad", "qquad", ",", ";", ":"}:
            return [_math_run(" ")]
        if command == "!":
            return []
        if command in _LATEX_MATH_SYMBOLS:
            return [_math_run(_LATEX_MATH_SYMBOLS[command])]
        return [_math_run(command)]

    return parse_sequence()


_LATEX_MATH_SYMBOLS = {
    "alpha": "\u03b1",
    "beta": "\u03b2",
    "gamma": "\u03b3",
    "Delta": "\u0394",
    "delta": "\u03b4",
    "theta": "\u03b8",
    "lambda": "\u03bb",
    "mu": "\u03bc",
    "nu": "\u03bd",
    "pi": "\u03c0",
    "rho": "\u03c1",
    "sigma": "\u03c3",
    "phi": "\u03c6",
    "omega": "\u03c9",
    "times": "\u00d7",
    "cdot": "\u00b7",
    "leq": "\u2264",
    "geq": "\u2265",
    "neq": "\u2260",
    "pm": "\u00b1",
    "infty": "\u221e",
    "rightarrow": "\u2192",
    "leftarrow": "\u2190",
}


def _math_run(text: str) -> Any:
    from docx.oxml import OxmlElement

    run = OxmlElement("m:r")
    text_element = OxmlElement("m:t")
    if text.startswith(" ") or text.endswith(" "):
        text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_element.text = text
    run.append(text_element)
    return run


def _math_arg(tag: str, children: list[Any] | None) -> Any:
    from docx.oxml import OxmlElement

    arg = OxmlElement(f"m:{tag}")
    for child in children or [_math_run("")]:
        arg.append(child)
    return arg


def _math_fraction(numerator: list[Any], denominator: list[Any]) -> Any:
    from docx.oxml import OxmlElement

    fraction = OxmlElement("m:f")
    fraction.append(_math_arg("num", numerator))
    fraction.append(_math_arg("den", denominator))
    return fraction


def _math_radical(children: list[Any]) -> Any:
    from docx.oxml import OxmlElement

    radical = OxmlElement("m:rad")
    radical.append(OxmlElement("m:deg"))
    radical.append(_math_arg("e", children))
    return radical


def _math_script(base: list[Any], subscript: list[Any] | None, superscript: list[Any] | None) -> Any:
    from docx.oxml import OxmlElement

    if subscript is not None and superscript is not None:
        element = OxmlElement("m:sSubSup")
        element.append(_math_arg("e", base))
        element.append(_math_arg("sub", subscript))
        element.append(_math_arg("sup", superscript))
        return element
    if subscript is not None:
        element = OxmlElement("m:sSub")
        element.append(_math_arg("e", base))
        element.append(_math_arg("sub", subscript))
        return element
    element = OxmlElement("m:sSup")
    element.append(_math_arg("e", base))
    element.append(_math_arg("sup", superscript))
    return element


def _linearize_latex(value: str) -> str:
    text = str(value or "").strip()
    for _ in range(6):
        next_text = re.sub(r"\\(?:d?frac|tfrac)\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", text)
        next_text = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", next_text)
        if next_text == text:
            break
        text = next_text
    symbols = {
        "alpha": "α", "beta": "β", "gamma": "γ", "Delta": "Δ", "delta": "δ",
        "theta": "θ", "lambda": "λ", "mu": "μ", "nu": "ν", "pi": "π",
        "rho": "ρ", "sigma": "σ", "phi": "φ", "omega": "ω", "times": "×",
        "cdot": "·", "leq": "≤", "geq": "≥", "neq": "≠", "pm": "±",
        "infty": "∞", "rightarrow": "→", "leftarrow": "←",
    }
    for command, symbol in symbols.items():
        text = re.sub(rf"\\{command}\b", symbol, text)
    text = re.sub(r"\\(?:text|mathrm|mathbf|operatorname)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:left|right|quad|qquad)\b|\\[!,;]", "", text)
    superscript = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
    subscript = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
    text = re.sub(r"\^\{([^{}]+)\}", lambda match: match.group(1).translate(superscript), text)
    text = re.sub(r"\^([0-9n+\-])", lambda match: match.group(1).translate(superscript), text)
    text = re.sub(r"_\{([^{}]+)\}", lambda match: match.group(1).translate(subscript), text)
    text = re.sub(r"_([0-9+\-])", lambda match: match.group(1).translate(subscript), text)
    text = re.sub(r"\\[A-Za-z]+", "", text)
    return text.replace("{", "").replace("}", "").strip()


def _format_run(run: Any, font: str, size: float, *, bold: bool = False, italic: bool = False, color: str | None = None) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def _docx_alignment(value: str) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
    }.get(value, WD_ALIGN_PARAGRAPH.CENTER)


def _word_font(value: str) -> str:
    return {
        "heiti": "SimHei",
        "kaiti": "KaiTi",
        "fangsong": "FangSong",
        "system": "Microsoft YaHei",
    }.get(value, "SimSun")


def _question_knowledge(question: dict[str, Any]) -> str:
    if question.get("knowledge_point"):
        return str(question["knowledge_point"])
    for point in question.get("knowledge_points") or []:
        if not isinstance(point, dict):
            continue
        for key in ("topic3_name", "topic2_name", "topic1_name"):
            if point.get(key):
                return str(point[key])
    return "课堂练习"


def _chunk_lines(lines: list[str], max_chars: int) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    length = 0
    for raw in lines:
        line = _office_text(str(raw or ""))
        if not line:
            continue
        pieces = [line[index : index + max_chars] for index in range(0, len(line), max_chars)] or [""]
        for piece in pieces:
            if current and length + len(piece) > max_chars:
                chunks.append(current)
                current = []
                length = 0
            current.append(piece)
            length += len(piece)
    if current:
        chunks.append(current)
    return chunks or [[""]]


def _bounded_number(value: Any, minimum: float, maximum: float, fallback: float) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        return fallback
    return max(minimum, min(maximum, parsed))


def _verify_office_package(path: Path, export_format: str) -> None:
    import zipfile

    if not path.is_file() or path.stat().st_size < 1000:
        raise ValueError("Generated Office file is empty")
    required = "word/document.xml" if export_format == "word" else "ppt/presentation.xml"
    with zipfile.ZipFile(path) as archive:
        if required not in archive.namelist():
            raise ValueError(f"Generated Office file is missing {required}")
        bad_member = archive.testzip()
        if bad_member:
            raise ValueError(f"Generated Office package is corrupt at {bad_member}")


def _portable_path(path: Path, root: Path) -> str:
    try:
        return path.resolve().relative_to(root).as_posix()
    except ValueError:
        return str(path.resolve())


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
